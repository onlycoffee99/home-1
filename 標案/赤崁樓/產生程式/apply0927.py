# -*- coding: utf-8 -*-
# 以使用者 1150927 版為底:補回修正 + 新增附件(履約實績佐證資料)
import re, sys, copy, docx
from docx.shared import Cm
from docx.oxml.ns import qn
from docx.oxml import parse_xml
SRC, OUT = sys.argv[1], sys.argv[2]
d = docx.Document(SRC); body = d.element.body
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
sid = {s.style_id: s for s in d.styles}
def sty(name): return next(s for s in d.styles if s.name == name)
def ptext(el): return ''.join(t.text or '' for t in el.iter(qn('w:t')))
def add_runs(p, text):
    for i, seg in enumerate(re.split(r'【(.*?)】', text)):
        if seg:
            r = p.add_run(seg)
            if i % 2: r.font.highlight_color = 7
def mk_par(style, text):
    p = d.add_paragraph(style=style); add_runs(p, text); body.remove(p._p); return p._p
def set_cell(cell, text):
    ts = list(cell._tc.iter(qn('w:t')))
    ts[0].text = text; ts[0].set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    for t in ts[1:]: t.text = ''
def set_par(el, text):
    ts = list(el.iter(qn('w:t'))); ts[0].text = text
    for t in ts[1:]: t.text = ''
def mk_figure(imgs, caption, per_row):
    total = 8646
    t = d.add_table(rows=0, cols=per_row); body.remove(t._tbl)
    tp = t._tbl.tblPr
    for old in tp.findall(qn('w:tblW')) + tp.findall(qn('w:tblStyle')): tp.remove(old)
    for x in (f'<w:tblStyle xmlns:w="{W}" w:val="a9"/>', f'<w:tblW xmlns:w="{W}" w:w="{total}" w:type="dxa"/>', f'<w:jc xmlns:w="{W}" w:val="center"/>', f'<w:tblLayout xmlns:w="{W}" w:type="fixed"/>'):
        tp.append(parse_xml(x))
    order = ['tblStyle', 'tblW', 'jc', 'tblLayout', 'tblLook']
    kids = sorted(list(tp), key=lambda e: order.index(e.tag.split('}')[1]) if e.tag.split('}')[1] in order else 99)
    for k in list(tp): tp.remove(k)
    for k in kids: tp.append(k)
    cw = total // per_row
    for gc in t._tbl.tblGrid.findall(qn('w:gridCol')): gc.set(qn('w:w'), str(cw))
    for i in range(0, len(imgs), per_row):
        cells = t.add_row().cells
        for j, (path, wcm) in enumerate(imgs[i:i + per_row]):
            c = cells[j]; tcPr = c._tc.get_or_add_tcPr()
            for old in tcPr.findall(qn('w:tcW')): tcPr.remove(old)
            tcPr.append(parse_xml(f'<w:tcW xmlns:w="{W}" w:w="{cw}" w:type="dxa"/>'))
            tcPr.append(parse_xml(f'<w:vAlign xmlns:w="{W}" w:val="center"/>'))
            p = c.paragraphs[0]; p.alignment = 1; p.add_run().add_picture(path, width=Cm(wcm))
    for tr in t._tbl.findall(qn('w:tr')):
        tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit xmlns:w="{W}"/>'))
        for p in tr.iter(qn('w:p')):
            ppr = p.get_or_add_pPr(); idx = 1 if len(ppr) and ppr[0].tag == qn('w:pStyle') else 0
            ppr.insert(idx, parse_xml(f'<w:keepNext xmlns:w="{W}"/>'))
    cap = t.add_row().cells
    m = cap[0].merge(cap[-1]) if per_row > 1 else cap[0]
    p = m.paragraphs[0]; p.style = sid['a']; add_runs(p, caption)
    ppr = p._p.get_or_add_pPr(); ppr.insert(1, parse_xml(f'<w:numPr xmlns:w="{W}"><w:ilvl w:val="0"/><w:numId w:val="0"/></w:numPr>'))
    return t._tbl

# ---- 1. 實績表修正 ----
tb = next(t for t in d.tables if t.rows[0].cells[0].text.strip() == '項次' and '類別' in t.rows[0].cells[1].text)
r1 = tb.rows[1].cells
set_cell(r1[2], '進駐輕車悠遊股份有限公司經營之宜蘭縣礁溪溫泉公園(營運移轉OT案)，於泡腳亭全區販賣部提供園區遊客餐飲服務，經營「時間到咖啡館－公園裡的咖啡館」；合約期間自民國112年2月1日至117年10月30日。')
set_cell(r1[3], '進駐合約書(附件1-1)')
r2 = tb.rows[2].cells
set_cell(r2[2], '時間到咖啡館台中館進駐臺中市第二市場停車場附屬空間(蚨聚有限公司受臺中市停車管理處委託經營之FU SPACE，與中區親子館共構)，租賃契約於民國115年3月20日經臺灣臺中地方法院所屬民間公證人公證(出租人：蚨聚有限公司；承租人：鑫和洋行)，於2026年9月19日開幕。')
set_cell(r2[3], '公證書、點交簽到表、開幕照片(附件1-2～1-4)')
set_cell(tb.rows[3].cells[3], '商標註冊證(圖5)')
set_cell(tb.rows[4].cells[3], '商品照片(附件1-5)')
set_cell(tb.rows[6].cells[3], '檢測報告(附件1-6)')
# 補第9項
if tb.rows[-1].cells[0].text.strip() != '9':
    new = copy.deepcopy(tb.rows[-1]._tr); tb._tbl.append(new)
    cells = tb.rows[-1].cells
    for c, v in zip(cells, ['9', '文化推廣課程', '2026年9月5日由鄭玉屏擔任講師，於臺中市西屯區國安一期社會住宅開設咖啡手沖體驗課程，帶領親子及社區民眾實作。', '課程照片(附件1-7)']):
        set_cell(c, v)

# ---- 2. 協力人員顧問角色(依 115/9/26 指示) ----
for t in d.tables:
    if t.rows[0].cells[0].text.strip() == '姓名' and '本案分工' in t.rows[0].cells[-1].text:
        for r in t.rows[1:]:
            n = r.cells[0].text
            if '蔡錦佳' in n: set_cell(r.cells[-1], '手作課程顧問：體驗課程規劃及教材設計建議')
            if '黃荻昌' in n: set_cell(r.cells[-1], '文化顧問：文化沙龍內容規劃')
for el in body.iter(qn('w:p')):
    s = ptext(el)
    if s.startswith('於本案擔任團隊之印刷手作課程顧問'):
        set_par(el, '於本案擔任團隊手作課程顧問，以其印刷工藝與視覺設計專業，可規劃「手作文創包裝」等體驗課程，設計課程流程、教材及適合親子與旅客之操作步驟，並指導團隊；課程作品可結合茶館文創展售區展示，延伸遊客之文化記憶。')
    if s.startswith('黃荻昌副教授任教於長榮大學') and '於本案擔任文化顧問' in s:
        set_par(el, s[:s.index('於本案擔任文化顧問')] + '於本案擔任文化顧問，可規劃「府城交通與城市記憶」系列文化沙龍內容。')

# ---- 3. 附件:履約實績佐證資料 ----
sect = body.find(qn('w:sectPr')); last = sect.getprevious()
els = [parse_xml(f'<w:p xmlns:w="{W}"><w:r><w:br w:type="page"/></w:r></w:p>'),
       mk_par(sty('A-01-壹'), '附件'),
       mk_par(sty('A-02-一'), '履約實績佐證資料'),
       mk_par(sty('B-02-一文'), '以下依「鑫和洋行近年業務實績一覽表」項次，檢附相關契約、公證文件及現場照片影本，以資佐證。')]
FIG = [
 ([('img/jx.jpg', 14)], '附件1-1　礁溪溫泉公園進駐合約書(甲方：輕車悠遊股份有限公司；乙方：鑫和洋行；合約期間112年2月1日至117年10月30日)', 1),
 ([('img/lease.jpg', 11)], '附件1-2　時間到咖啡館台中館租賃契約公證書(115年3月20日，臺灣臺中地方法院所屬民間公證人)', 1),
 ([('img/handover.jpg', 11)], '附件1-3　臺中市第二市場停車場附屬空間委託經營案點交簽到表(114年12月22日)', 1),
 ([('img/open1.jpg', 7.4), ('img/open2.jpg', 7.4)], '附件1-4　時間到咖啡館台中館開幕活動(2026年9月19日，臺中市停車管理處、中區區公所出席)', 2),
 ([('img/dripbag.jpg', 5.4), ('img/kumquat.jpg', 5.4), ('img/callig_pack.jpg', 3.9)], '附件1-5　自有品牌商品：舞荳掛耳咖啡、宜蘭金棗黑糖金磚、舞荳單品咖啡包裝', 3),
 ([('img/test1.jpg', 5.8), ('img/test2.jpg', 5.8)], '附件1-6　委託財團法人金屬工業研究發展中心產品成分分析報告(2019年3月)', 2),
 ([('img/course1.jpg', 7.4), ('img/course2.jpg', 7.4)], '附件1-7　國安一期社會住宅咖啡手沖體驗課程(2026年9月5日，講師鄭玉屏)', 2),
]
for imgs, cap, n in FIG:
    els.append(mk_figure(imgs, cap, n))
    els.append(mk_par(sid['af4'], '資料來源：本計畫彙整'))
for e in els:
    last.addnext(e); last = e

for i, dp in enumerate(body.iter('{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr'), 1):
    dp.set('id', str(i))
d.save(OUT); print('saved')
