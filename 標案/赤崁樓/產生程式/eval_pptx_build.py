import re,copy
from pptx import Presentation
from pptx.util import Inches,Pt,Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
W='/tmp/claude-0/-home-user-home-1/c0e478eb-8463-5374-8afd-baba07237450/scratchpad/eval/'
C=W+'cov/';IM=W+'deck/'
F='微軟正黑體'
RED=RGBColor(0xB9,0x4A,0x37);INK=RGBColor(0x3B,0x2A,0x20);MUT=RGBColor(0x5E,0x4E,0x40)
PAP=RGBColor(0xFB,0xF7,0xF0);PAP2=RGBColor(0xF1,0xE7,0xD8);BOR=RGBColor(0xD9,0xC9,0xB0);WHITE=RGBColor(0xFF,0xFF,0xFF);DARK=RGBColor(0x4A,0x2E,0x24)

def notes_of(i):
    s=open(f'{IM}project/slides/{i}.html').read()
    return re.search(r'<aside>(.*?)</aside>',s,re.S).group(1)

prs=Presentation(C+'c.pptx')
# drop slide 3 (placeholder 哭啊)
sl=prs.slides._sldIdLst; rid=sl[2].rId; prs.part.drop_rel(rid); sl.remove(sl[2])
BLANK=[l for l in prs.slide_layouts if l.name=='空白'][0]
SW,SH=prs.slide_width,prs.slide_height

def run_fmt(r,size,color=INK,bold=False):
    r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=color;r.font.name=F
    rPr=r._r.get_or_add_rPr()
    for tag in ('a:ea','a:latin'):
        pass
    from pptx.oxml.ns import qn
    ea=rPr.find(qn('a:ea'))
    if ea is None:
        ea=rPr.makeelement(qn('a:ea'),{});rPr.append(ea)
    ea.set('typeface',F)

def text(slide,x,y,w,h,paras,size=16,color=INK,bold=False,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,ls=1.25,bullet=False):
    tb=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=tb.text_frame;tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(0.05);tf.margin_top=tf.margin_bottom=Inches(0.03);tf.vertical_anchor=anchor
    if isinstance(paras,str): paras=[paras]
    for k,pt in enumerate(paras):
        p=tf.paragraphs[0] if k==0 else tf.add_paragraph();p.alignment=align;p.line_spacing=ls
        if bullet: p.space_after=Pt(6)
        segs=pt if isinstance(pt,list) else [(('• ' if bullet else '')+pt,bold,color)]
        for sgt in segs:
            t,b,c=(sgt+(color,))[:3] if len(sgt)==2 else sgt
            r=p.add_run();r.text=t;run_fmt(r,size,c,b)
    return tb

def rect(slide,x,y,w,h,fill,line=None,rounded=True):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    if rounded: sh.adjustments[0]=0.06
    sh.fill.solid();sh.fill.fore_color.rgb=fill
    if line: sh.line.color.rgb=line;sh.line.width=Pt(1)
    else: sh.line.fill.background()
    sh.shadow.inherit=False
    return sh

def card(slide,x,y,w,h,title,body,fill=PAP,tc=RED,bc=MUT,ts=20,bs=15):
    rect(slide,x,y,w,h,fill,BOR if fill==PAP else None)
    text(slide,x+0.2,y+0.15,w-0.4,0.5,title,ts,tc,True)
    if body: text(slide,x+0.2,y+0.7,w-0.4,h-0.8,body,bs,bc,ls=1.35)

def pic(slide,path,x,y,w,h):
    im=Image.open(path);iw,ih=im.size;r=w/h;ir=iw/ih
    p=slide.shapes.add_picture(path,Inches(x),Inches(y),Inches(w),Inches(h))
    if ir>r: c=(1-r/ir)/2;p.crop_left=p.crop_right=c
    else: c=(1-ir/r)/2;p.crop_top=p.crop_bottom=c
    return p

def table(slide,x,y,w,head,rows,cw,fs=14,rh=0.42):
    n=len(rows)+1;sh=slide.shapes.add_table(n,len(head),Inches(x),Inches(y),Inches(w),Inches(rh*n));t=sh.table
    tblPr=sh._element.graphic.graphicData.tbl.tblPr;tblPr.set('bandRow','0');tblPr.set('firstRow','0')
    st=tblPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId')
    if st is not None: st.text='{5940675A-B579-460E-94D1-54222C63F5DA}'
    for j,f in enumerate(cw): t.columns[j].width=Inches(w*f)
    for i in range(n):
        t.rows[i].height=Inches(rh)
        for j in range(len(head)):
            c=t.cell(i,j);c.margin_left=c.margin_right=Inches(0.1);c.margin_top=c.margin_bottom=Inches(0.04);c.vertical_anchor=MSO_ANCHOR.MIDDLE
            v=head[j] if i==0 else rows[i-1][j];b=False
            if isinstance(v,tuple): v,b=v
            c.fill.solid();c.fill.fore_color.rgb=RED if i==0 else (PAP if i%2 else PAP2)
            tf=c.text_frame;tf.word_wrap=True;p=tf.paragraphs[0];r=p.add_run();r.text=v
            run_fmt(r,fs,WHITE if i==0 else INK,i==0 or b)
    return sh

page=[2]
def new(eyebrow,title,nid,notes=None):
    s=prs.slides.add_slide(BLANK)
    s.shapes.add_picture(C+'content_bg.jpg',0,0,SW,SH)
    rect(s,0,0,0.18,7.5,RED,rounded=False)
    text(s,0.6,0.32,11,0.35,eyebrow,13,RED,True)
    text(s,0.6,0.62,12,0.7,title,28,INK,True)
    ln=s.shapes.add_connector(1,Inches(0.65),Inches(1.4),Inches(12.7),Inches(1.4));ln.line.color.rgb=BOR;ln.line.width=Pt(1.25)
    page[0]+=1
    text(s,0.6,7.0,8,0.3,'鑫和洋行｜赤崁樓遊客中心2樓部分空間委外經營管理案',10,MUT)
    text(s,12.0,7.0,0.8,0.3,f'{page[0]:02d}',10,MUT,align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text=notes or notes_of(nid)
    return s

E1='評選項目一｜專業能力與績效（10分）';E2='評選項目二｜經營管理內容、創意能力及經營策略（25分）';E3='評選項目三｜權利金報價合理性（20分）';E4='評選項目四｜財務規劃（15分）';E5='評選項目五｜空間使用規劃（15分）';E6='評選項目六｜環境、建物及設備之管理維護計畫（15分）'

# --- cover (slide1)
s1=prs.slides[0]
for sh in s1.shapes:
    if sh.has_text_frame and '115年' in sh.text_frame.text:
        sh.width=Inches(5.2);sh.text_frame.word_wrap=False;sh.left=Emu(1524000);sh.top=Inches(6.05)
        for p in sh.text_frame.paragraphs:
            for r in p.runs: r.text=r.text.replace('XX','10')
    if sh.has_text_frame and '投標廠商' in sh.text_frame.text:
        r=sh.text_frame.paragraphs[0].runs
        tf=sh.text_frame;p=tf.add_paragraph()
        pPr0=tf.paragraphs[0]._p.find('{http://schemas.openxmlformats.org/drawingml/2006/main}pPr')
        if pPr0 is not None: p._p.insert(0,copy.deepcopy(pPr0))
        rr=p.add_run();rr.text='計畫主持人∣鄭玉屏'
        sh.height=Inches(0.9)
        rr.font.bold=r[0].font.bold;rr.font.size=r[0].font.size
        rPr=r[0]._r.find('{http://schemas.openxmlformats.org/drawingml/2006/main}rPr')
        if rPr is not None: rr._r.insert(0,copy.deepcopy(rPr))
s1.notes_slide.notes_text_frame.text=notes_of('cover')

# --- 目錄 (slide2)
s2=prs.slides[1]
items=[('壹','專業能力與績效','10分'),('貳','經營管理內容、創意能力及經營策略','25分'),('參','權利金報價合理性','20分'),('肆','財務規劃','15分'),('伍','空間使用規劃','15分'),('陸','環境、建物及設備之管理維護計畫','15分')]
# remove template placeholders except title
for sh in list(s2.shapes):
    if sh.is_placeholder and sh.placeholder_format.type!=None and 'TITLE' not in str(sh.placeholder_format.type):
        sh._element.getparent().remove(sh._element)
s2._element.set('showMasterSp','0')
bgp=s2.shapes.add_picture(C+'toc_bg.jpg',0,0,SW,SH);s2.shapes._spTree.remove(bgp._element);s2.shapes._spTree.insert(2,bgp._element)
ar=s2.shapes.add_shape(MSO_SHAPE.PENTAGON,Emu(0),Emu(615951),Emu(2324100),Emu(982662));ar.fill.solid();ar.fill.fore_color.rgb=RGBColor(0xD8,0xC2,0xA8);ar.line.fill.background();ar.shadow.inherit=False
tsp=[x for x in s2.shapes if x.is_placeholder][0]._element;tsp.getparent().remove(tsp);s2.shapes._spTree.append(tsp)
y0=1.75
for k,(a,b,c) in enumerate(items):
    y=y0+k*0.82
    o=s2.shapes.add_shape(MSO_SHAPE.OVAL,Inches(3.05),Inches(y),Inches(0.62),Inches(0.62));o.fill.solid();o.fill.fore_color.rgb=RGBColor(0xEA,0xD8,0xC6);o.line.fill.background()
    tf=o.text_frame;tf.margin_left=tf.margin_right=0;p=tf.paragraphs[0];p.alignment=PP_ALIGN.CENTER;r=p.add_run();r.text=a;run_fmt(r,18,RED,True)
    text(s2,3.85,y+0.08,6.2,0.5,[[(b,True,INK),('　'+c,False,MUT)]],18)
s2.notes_slide.notes_text_frame.text=notes_of('agenda')

# --- 一
s=new(E1,'鑫和洋行：公共場域經營實績','record')
table(s,0.6,1.65,8.3,['類別','實績內容'],[
 ['公共場域進駐','宜蘭礁溪溫泉公園「時間到咖啡館－公園裡的咖啡館」'],
 ['跨區門市','台中館：承租蚨聚有限公司經營之臺中市第二市場停車場附屬空間，115/9/19開幕'],
 ['品牌資產','「舞荳及圖」商標註冊第02507640號'],
 ['產學合作','佛光大學「玩咖系列」濾掛咖啡；USR國際交流'],
 ['產品檢測','2019年起委託金屬中心成分分析'],
 ['推廣教學','國安一期社宅咖啡手沖課（講師鄭玉屏）']],[0.2,0.8],fs=14,rh=0.66)
pic(s,IM+'opening.jpg',9.2,1.65,3.5,2.3);text(s,9.2,3.98,3.5,0.6,'台中館開幕，臺中市停車管理處、中區區公所出席',11,MUT)
pic(s,IM+'jiaoxi.jpg',9.2,4.6,3.5,2.05);text(s,9.2,6.66,3.5,0.3,'礁溪溫泉公園 公園裡的咖啡館',11,MUT)

s=new(E1,'計畫主持人與承攬緣由','pi')
text(s,0.6,1.75,7.6,2.2,['鄭玉屏：鑫和洋行營運長、本案主理人，咖啡烘焙師','店面SOP、菜單設計、品項開發與人員訓練','協辦宜蘭冬戀溫泉季、佛光大學文創市集'],18,MUT,bullet=True)
rect(s,0.6,4.1,7.6,1.9,PAP,BOR);rect(s,0.6,4.1,0.1,1.9,RED,rounded=False)
text(s,0.9,4.25,7.1,1.7,[[('承攬緣由：',True,RED),('「舞荳」品牌源自陳玨霖老師題贈墨寶，我們一直在做「把餐飲和文化變成生活產品」；赤崁樓是最適合實踐這件事的場域。',False,INK)]],17,ls=1.4)
pic(s,IM+'calligraphy.jpg',8.8,1.65,3.9,4.9)

s=new(E1,'協力團隊：由鑫和洋行單獨投標','team')
table(s,0.6,1.7,12.1,['協力人員／單位','背景','本案分工'],[
 ['蔡錦佳老師','東華印刷局資訊科技企業社計畫主持人','印刷手作課程顧問'],
 ['黃荻昌副教授','長榮大學副教授，鐵道工程與文化資產','文化顧問、文化沙龍規劃'],
 ['徐鈺清','台灣瑰寶牛樟芝生物科技負責人','原生物產顧問、泡茶用水研討'],
 ['廣廓生活有限公司','Himalaya Smile 天然手作商品','文創選物規劃'],
 ['傳宇智聯有限公司','AI企業培訓及品牌行銷','社群經營與AI行銷']],[0.24,0.42,0.34],fs=16,rh=0.6)
text(s,0.6,5.55,12,0.5,'均已簽署合作意向書（附件一）。本案依投標須知由鑫和洋行單獨投標，負全部履約責任。',15,MUT)

# --- 二 vision statement on toc bg
s=prs.slides.add_slide(BLANK);s.shapes.add_picture(C+'toc_bg.jpg',0,0,SW,SH);page[0]+=1
text(s,0.9,1.3,8,0.4,E2,14,RED,True)
text(s,0.9,2.0,8.5,2.2,['守護場域、服務遊客、','轉譯文化'],48,RED,True,ls=1.15)
text(s,0.9,4.4,7.6,1.4,'經營目標：讓赤崁樓遊客中心2樓成為兼具文化深度、服務品質與營運韌性的文化複合空間。',18,INK,ls=1.45)
text(s,12.0,7.0,0.8,0.3,f'{page[0]:02d}',10,MUT,align=PP_ALIGN.RIGHT)
s.notes_slide.notes_text_frame.text=notes_of('vision')

s=new(E2,'經營內容：老茶新喝，臺灣茶手搖飲為主力','product')
text(s,0.6,1.75,5.2,3.2,['符合工作說明書「以輕飲食為規劃主軸」。','傳統好茶以原茶、奶茶、果茶、茶凍、茶冰沙重新詮釋；每款茶附產區與製茶故事卡。'],18,INK,ls=1.45)
text(s,0.6,4.6,5.2,1.5,'多語菜單（中英日韓）、環保杯折扣、兒童杯與無咖啡因選項。',15,MUT,ls=1.4)
table(s,6.3,1.65,6.4,['收費項目','收費標準（元）'],[['臺灣原茶','50～90'],['奶茶與茶拿鐵','65～110'],['季節果茶（在地農產）','60～120'],['茶凍與茶冰沙','60～110'],['品茶組合','150～250'],['舞荳咖啡','90～180'],['甜品與茶點','50～150'],['文化體驗','免費～400／人']],[0.6,0.4],fs=15,rh=0.52)

s=new(E2,'創意能力：文化內容與營運互相支持','culture')
table(s,0.6,1.75,12.1,['體驗','現場做法','對應收入'],[
 ['一杯茶認識臺灣','圖文介紹茶的製作與外銷故事','單杯茶飲、品茶組合'],
 ['赤崁建築與故事','圖卡呈現普羅民遮堡、贔屭馱碑','延長停留，帶動飲品甜品'],
 ['夏季也能喝茶','同一茶款延伸為茶凍與茶冰沙','不同季節與客群銷售'],
 ['智慧供茶','多語選茶、付款與標準化出茶','提升尖峰服務效率'],
 ['故事活動','茶文化體驗、親子故事場、團體預約','活動收入與團體消費']],[0.24,0.44,0.32],fs=17,rh=0.68)

s=new(E2,'行銷推廣：四季主題，串聯在地','season')
for k,(a,b) in enumerate([('春季','春茶品飲會；兒童節親子故事場及茶點手作'),('夏季','冷泡茶與茶冰沙；暑期學生文化體驗營；鐵道文化講座'),('秋季','古蹟月文化沙龍：赤崁樓建築變遷與贔屭馱碑導覽'),('冬季','年終選物市集、活版印刷賀卡手作；新春茶禮盒')]):
    card(s,0.6+k*3.06,1.75,2.86,2.9,a,b,ts=22,bs=15)
text(s,0.6,5.0,12.1,1.2,'IG、Facebook、LINE官方帳號、Google商家每週至少3則貼文；串聯赤嵌朋派商圈、赤嵌社區等在地組織；宣傳內容事先報機關同意。',15,MUT,ls=1.45)

s=new(E2,'組織與人力：平日3人、假日5人在場','staff')
table(s,0.6,1.7,8.0,['職稱','人數','工作職掌'],[['計畫主持人','1','鄭玉屏，統籌履約與機關聯繫'],['店長','1','現場管理、食安及消防安全管理人'],['茶飲調製人員','3','茶飲及咖啡製作、商品解說'],['外場服務人員','1','接待點餐、環境整理、遊客諮詢'],['兼職人員','2～3','假日及活動支援，優先聘在地居民'],['行政及會計','1（兼）','帳務、採購、報表、保險']],[0.24,0.13,0.63],fs=14,rh=0.6)
card(s,8.95,1.7,3.75,3.4,'教育訓練','職前訓練16小時（古蹟規範、食安、消防、赤崁樓導覽）；每年食品衛生講習8小時；每季服務及茶飲專業訓練',ts=20,bs=15)

s=new(E2,'執行進度：配合遊客中心完工點交','schedule','（約30秒）遊客中心目前仍在興建，預計約兩年後完工點交。我們把這段時間當作籌備期：決標後14日內簽約，30日內提送營運管理維護計畫；工程期間在礁溪與台中館先研發、試賣赤崁主題茶款，建立在地合作與社群預熱，並完成人員培訓。完工點交後依竣工現況修正設計圖說送審，核定後施作可復原裝修，目標點交後2個月內開幕，開幕時商品與客群都已就緒。')
steps=[('決標簽約','決標後14日內'),('提送計畫','簽約後30日內提送營運管理維護計畫及初步設計'),('籌備期','遊客中心興建期間（約2年）：茶款研發試賣、在地合作、預熱行銷、人員培訓'),('完工點交','依機關通知會同清點；依竣工現況修正設計圖說送審'),('裝修開幕','核定後施作可復原裝修；目標點交後2個月內開幕，營運5年')]
for k,(a,b) in enumerate(steps):
    x=0.6+k*2.44
    ch=s.shapes.add_shape(MSO_SHAPE.CHEVRON if k else MSO_SHAPE.PENTAGON,Inches(x),Inches(1.8),Inches(2.5),Inches(0.85));ch.fill.solid();ch.fill.fore_color.rgb=RED if k==2 else DARK;ch.line.fill.background();ch.shadow.inherit=False
    tf=ch.text_frame;p=tf.paragraphs[0];p.alignment=PP_ALIGN.CENTER;r=p.add_run();r.text=a;run_fmt(r,18,WHITE,True)
    rect(s,x+0.05,2.85,2.25,2.6,PAP if k!=2 else RGBColor(0xF6,0xE2,0xDA),BOR)
    text(s,x+0.18,2.98,2.0,2.4,b,15,INK if k==2 else MUT,k==2,ls=1.4)
text(s,0.6,5.75,12.1,0.9,'籌備期即在礁溪、台中館試賣赤崁主題茶款並累積社群，開幕時商品與客群已就緒；審查期間依契約得申請營運期展延。',15,MUT,ls=1.4)

s=new(E2,'執行能力：對應年度考核五大項目','review')
table(s,0.6,1.7,12.1,['考核項目（權重）','本團隊做法'],[
 ['營運計畫執行（含財務）25%','依核定計畫營運；會計獨立、每月損益檢討、營運月報'],
 ['設備維修及環境清潔 25%','設備保養表定執行並建檔；每2小時巡檢公共區域'],
 ['歷史、文化及藝術教育推廣 25%','每年24場以上文化沙龍及體驗；茶款附故事卡'],
 ['與文化局配合度 15%','機關借用場地優先無償提供；宣傳事先報准；主持人親自出席考核'],
 ['顧客滿意度或投訴率 10%','每半年滿意度調查，目標85%以上；意見3日內回覆']],[0.34,0.66],fs=15,rh=0.62)
text(s,0.6,5.6,12,0.5,'自我評估目標：年度考核85分以上（75分通過），作為後續擴充3年之依據。',15,MUT)

# --- 三
s=new(E3,'固定權利金：依周邊租金行情推估','basis')
table(s,0.6,1.75,6.6,['周邊租金熱點','一樓每坪月租金（元）'],[['中正路、國華街三段','2,100～3,273'],['中山路、西門路二段','889～3,375'],['民族路二段（參考基準）','1,200～1,472'],['民生路一段','596～833']],[0.56,0.44],fs=16,rh=0.62)
text(s,0.6,4.95,6.6,0.4,'資料來源：591房屋交易網，本團隊彙整分析',12,MUT)
rect(s,7.6,1.75,5.1,4.0,DARK)
text(s,7.85,1.95,4.6,1.6,['1,300元 × 70%（商圈繁榮度）× 75%（二樓樓層效用）＝ 683元','（683 ＋ 管理費62元）× 12月 × 61.51坪'],15,RGBColor(0xEA,0xD8,0xC6),ls=1.4)
text(s,7.85,3.65,4.6,0.8,'≈ 549,899元／年',30,WHITE,True)
text(s,7.85,4.6,4.6,0.9,'報價：每年55萬元，高於契約底價54萬元',16,RGBColor(0xF2,0xB8,0xA6),True)

s=new(E3,'固定保底，營收超過目標值再分享','royalty')
rect(s,0.6,1.7,5.95,2.0,DARK);text(s,0.85,1.82,5.5,0.4,'固定權利金（外加5%營業稅）',14,RGBColor(0xEA,0xD8,0xC6));text(s,0.85,2.2,5.5,0.7,'每年55萬元',32,WHITE,True);text(s,0.85,3.0,5.5,0.5,'不論營運好壞照繳；五年合計275萬元',14,RGBColor(0xEA,0xD8,0xC6))
rect(s,6.75,1.7,5.95,2.0,PAP,BOR);text(s,7.0,1.82,5.5,0.4,'變動權利金（外加5%營業稅）',14,MUT);text(s,7.0,2.2,5.5,0.7,'超過540萬部分×2%',32,RED,True);text(s,7.0,3.0,5.5,0.5,'契約第4條：未達540萬元免繳；依401申報書計算',14,MUT)
table(s,0.6,3.95,12.1,['年度','預估年營業額','超過540萬部分','變動權利金（2%）','權利金合計'],[
 ['第一年','8,327,040','2,927,040','58,541',('608,541',True)],['第二年','9,576,096','4,176,096','83,522',('633,522',True)],['第三年','10,096,536','4,696,536','93,931',('643,931',True)]],[0.14,0.22,0.22,0.22,0.20],fs=16,rh=0.52)
text(s,0.6,6.15,12,0.4,'單位：新臺幣元（未稅）；營業額依服務建議書表15。',12,MUT)

# --- 四
s=new(E4,'投資金額與經費來源','invest')
table(s,0.6,1.7,6.9,['項目','金額（元）'],[['可復原式室內裝修及木作','800,000'],['吧台及餐飲設備','550,000'],['家具及陳設','250,000'],['資訊及安全設備','100,000'],[('開辦小計',True),('1,700,000',True)],['營運週轉金（約3個月）','300,000'],[('總投入資金',True),('2,000,000',True)]],[0.64,0.36],fs=15,rh=0.6)
rect(s,7.9,1.7,4.8,1.45,DARK);text(s,8.15,1.8,4.3,0.4,'自有資金',14,RGBColor(0xEA,0xD8,0xC6));text(s,8.15,2.2,4.3,0.7,'140萬元（70%）',30,WHITE,True)
rect(s,7.9,3.35,4.8,1.45,PAP,BOR);text(s,8.15,3.45,4.3,0.4,'金融機構貸款',14,MUT);text(s,8.15,3.85,4.3,0.7,'60萬元（30%）',30,RED,True)
text(s,7.9,5.05,4.8,1.2,'開辦費用分3年攤提；每年另自盈餘提撥營業額1%為設備汰換準備金',14,MUT,ls=1.4)

s=new(E4,'分年收支預算：以來客數達基本面為前提','finance','（約30秒）財務採保守估算，營運年度自點交次日起算。第一年只估穩定水準的八成，約損益兩平；第二、三年稅前淨利率約6.6%與8.1%。前提是來客數達到基本面：每日平均約207人次才能損益兩平。點交前將依當時物價隨營運管理維護計畫更新財務計畫。地價稅、房屋稅由廠商負擔，已列在雜支。')
table(s,0.6,1.65,8.6,['項目（元）','第一年','第二年','第三年'],[['營業收入','8,327,040','9,576,096','10,096,536'],['原物料、包材及商品','3,078,125','3,539,844','3,732,227'],['人事費用','2,880,000','2,966,400','3,055,392'],['權利金（固定＋變動）','608,541','633,522','643,931'],['水電、行銷、保險、雜支等','1,141,258','1,235,046','1,281,853'],['開辦攤提','566,667','566,667','566,667'],[('稅前損益',True),('52,449',True),('634,617',True),('816,466',True)]],[0.37,0.21,0.21,0.21],fs=14,rh=0.52)
rect(s,9.5,1.65,3.2,3.0,RED);text(s,9.7,1.8,2.8,0.4,'損益兩平',14,RGBColor(0xFB,0xE3,0xDA));text(s,9.7,2.25,2.8,0.8,'207人次',34,WHITE,True);text(s,9.7,3.15,2.8,1.4,'每日平均來客；穩定水準估260人次、客單價115元',13,RGBColor(0xFB,0xE3,0xDA),ls=1.35)
text(s,0.6,5.95,12.1,0.8,'營運年度自點交次日起算；營收爬坡80%／92%／97%；稅前淨利率0.6%／6.6%／8.1%。地價稅、房屋稅列於雜支。',14,MUT)

# --- 五
s=new(E5,'空間使用規劃：約62坪，動線清楚','space','（約20秒）委外範圍203.35平方公尺、約62坪，分為六區。吧台與洗滌設備靠北側，鄰近既有給排水；文化沙龍區平日為座位，活動時轉為講座空間，可容納約30人。實際配置於完工點交時丈量後調整。')
rect(s,0.6,1.7,6.9,3.6,WHITE,BOR,rounded=False)
pp=s.shapes.add_picture(IM+'plan.png',Inches(0.7),Inches(1.8),Inches(6.7),Inches(3.35))
text(s,0.6,5.4,6.9,0.4,'紅框為委外範圍（招標文件平面圖）；實際配置依完工點交丈量調整',12,MUT)
table(s,7.9,1.7,4.8,['分區','面積'],[['入口接待及展售區','8坪'],['手搖飲吧台及備料區','12坪'],['休憩座位區（約36席）','20坪'],['文化沙龍多功能區（約30人）','14坪'],['倉儲及員工區','4坪'],['公共動線（淨寬1.2m以上）','4坪']],[0.76,0.24],fs=14,rh=0.55)

s=new(E5,'裝修設計：不破壞古蹟、可復原','design','（約20秒）裝修以可復原為原則：活動式家具、不鑽孔、不變更結構；吧台全部電熱，不用明火。因遊客中心尚在興建，簽約後先提初步設計，完工點交後再依竣工現況修正圖說送審，核定後才施作。')
pic(s,IM+'teahouse.jpg',0.6,1.7,6.9,4.3)
text(s,0.6,6.05,6.9,0.4,'空間意象示意圖，實際以機關核定設計圖說為準',12,MUT)
text(s,7.9,1.8,4.8,4.4,['活動式家具與模組化櫃體，不鑽孔、不變更結構','吧台全部電熱設備，不用明火及桶裝瓦斯','茶葉綠、原木與暖光；照明避免直射古蹟視野','無障礙座位與通道','簽約30日內提初步設計；完工點交後依竣工現況修正送審，核定後施作'],16,MUT,bullet=True,ls=1.35)

s=new(E5,'開放時間：優於契約要求','ops')
for k,(big,cap) in enumerate([('6日','每週營業（週二至週日）\n契約要求不少於5日'),('9小時','每日10:00–19:00\n契約要求不少於8小時'),('假日照常','國定假日、連續假期照常營業')]):
    x=0.6+k*4.08;rect(s,x,1.75,3.88,2.7,PAP,BOR);text(s,x+0.25,1.95,3.4,1.0,big,40,RED,True);text(s,x+0.25,3.1,3.4,1.2,cap.split('\n'),15,MUT,ls=1.35)
text(s,0.6,4.8,12.1,1.2,'週一為公休及設備保養日；開放時間於現場明顯處設告示牌；配合機關活動可延長；調整均事先取得機關書面同意。',15,MUT,ls=1.45)

# --- 六
s=new(E6,'環境、建物及設備管理維護','manage')
for k,(a,b) in enumerate([('清潔與設備','每日開閉店清潔、每2小時巡檢；空調、給排水、冷藏冷凍等依表定保養並登記建檔；聯網設備不用中國品牌'),('防災與防盜','設防火管理人；每半年至少1次防災演練並邀機關參與；監視錄影與夜間保全；現金每日存入銀行'),('保險','依契約額度投保公共意外責任險、火災險（250萬元以上）及產品責任險，保單送機關備查')]):
    card(s,0.6+k*4.08,1.75,3.88,3.3,a,b,ts=22,bs=15)
text(s,0.6,5.35,12.1,1.0,'點交時建立財產及設備清冊並拍照存檔；損壞即時通報機關並負責修復；期滿15日內點交返還。',15,MUT,ls=1.45)

# --- closing on cover bg
s=prs.slides.add_slide(BLANK);s.shapes.add_picture(C+'cover_bg.jpg',0,0,SW,SH);page[0]+=1
text(s,1.6,1.4,8,2.0,['讓赤崁樓的故事，','留在旅客手中的一杯茶裡'],36,RED,True,ls=1.25)
text(s,1.6,3.6,6,0.6,'鑫和洋行　敬請指教',22,INK,True)
s.notes_slide.notes_text_frame.text=notes_of('closing')

prs.save(W+'赤崁樓評選簡報_1151001.pptx');print('slides',len(prs.slides))
