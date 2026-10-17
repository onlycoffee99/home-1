# -*- coding: utf-8 -*-
"""把評選講稿 md 逐頁寫進簡報備忘稿(115/10/14)
用法:python3 -I notes_from_script.py 講稿.md 簡報.pptx 輸出.pptx
"""
import re
import sys

from pptx import Presentation

md, src, dst = sys.argv[1:4]
body = open(md, encoding='utf-8').read().split('## 二、時間對照')[0]
notes = {}
for b in re.split(r'\n(?=\*\*第\d)', body):
    m = re.match(r'\*\*第(\d+)(?:~(\d+))?頁　([^*]*)\*\*\(([^)]*)\)\n?(.*)', b, re.S)
    if not m:
        continue
    a, z = int(m.group(1)), int(m.group(2) or m.group(1))
    txt = re.sub(r'\n---.*', '', m.group(5).strip(), flags=re.S).strip().replace('*', '')
    if a == 16:
        for n, t in re.findall(r'- 第(\d+)頁:(.*)', txt):
            notes[int(n)] = f'(約7秒)\n{t}'
        notes[20] += '\n收尾:顧問負責規劃指導,現場一律由我們的人執行。\n10分鐘版:16~20頁連翻,只講「五位協力夥伴分別負責手作課程、文化沙龍、原生物產、文創選物、AI行銷,現場由本團隊執行」。'
        continue
    for k in range(a, z + 1):
        head = f'【第{a}{"~" + str(z) if z != a else ""}頁｜{m.group(4)}】\n'
        notes[k] = head + (txt if k == a else '(續上頁,邊講邊換頁)')
p = Presentation(src)
for i, s in enumerate(p.slides, 1):
    s.notes_slide.notes_text_frame.text = notes.get(i, '(章節頁,直接換頁)')
p.save(dst)
print(len(notes))
