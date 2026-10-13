# -*- coding: utf-8 -*-
"""陳玨霖老師(本公司品牌文化顧問)單頁投影片(115/10/13)
用法:python3 -I chen_slide.py 團隊組織_1151013.pptx 輸出.pptx
以協力人員備用頁(蔡錦佳)為版型,刪除其他頁,只留一頁;插入完整簡報第11頁(計畫主持人)之後。
資料來源:資歷庫/協力夥伴.md(台灣華夏文教農經交流協會網站、葛瑪蘭新聞網 2022-02-27)
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from team_slides_fill import BROWN, DARK, RED, T, find, place, remove, set_text  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Cm  # noqa: E402


def one_line(shp, text):
    runs = shp.text_frame.paragraphs[0].runs
    runs[0].text = text
    for r in runs[1:]:
        r.text = ''


def main(src, dst):
    prs = Presentation(src)
    keep = [s for s in prs.slides if any(sh.has_text_frame and sh.text_frame.text.strip() == '蔡錦佳' for sh in s.shapes)][-1]
    sld = prs.slides._sldIdLst
    for s in list(prs.slides):
        if s is not keep:
            el = [e for e in sld if int(e.get('id')) == s.slide_id][0]
            prs.part.drop_rel(el.rId)
            sld.remove(el)
    s = keep
    one_line(find(s, '文字方塊 1'), '陳玨霖')
    role = find(s, '文字方塊 17')
    place(role, w=16)
    set_text(role, [[T('品牌文化顧問(本公司)', 28, True, BROWN)]])
    set_text(find(s, '文字方塊 21'), [[T('宜蘭縣書法學會創會理事長 •書法教學近五十年 •「舞荳」品牌題字', 20)]])
    for th, td, head, desc in [('文字方塊 40', '文字方塊 45', '書法五體', '篆、隸、楷、行、草'),
                               ('文字方塊 70', '文字方塊 72', '書法教學', '社大講師、種子教師'),
                               ('文字方塊 75', '文字方塊 77', '跨界展演', '水墨舞蹈現場揮毫')]:
        one_line(find(s, th), head)
        set_text(find(s, td), [[T(desc, 16)]], wrap=False, algn='just')
    cv = [('現職　', '宜蘭縣書法學會指導教授、宜蘭縣社區大學書法講師、華夏書化研究學會理事長'),
          ('經歷　', '1978年起於宜蘭從事書法教學；1991年中國書法學會宜蘭縣支會首屆理事長；宜蘭縣書法學會創會理事長'),
          ('師承　', '康灩泉、吳斌遊等書畫名家'),
          ('展演　', '2013北京國際文化藝術交流大會書法現場表演特別金獎；2012「陳玨霖六十書法回顧展」；日本、巴黎展出')]
    body = find(s, '文字方塊 79')
    set_text(body, [[T('■ ', 11, False, RED), T(h, 15, True, BROWN), T(t, 15, False, BROWN)] for h, t in cv], line=110)
    set_text(find(s, 'TextBox 11'),
             [[T('與本公司關係｜品牌文化顧問：', 18, True, RED),
               T('「舞荳」二字由老師親筆題贈(2024年)，提供品牌書法與文化指導；本案文化內容由協力人員規劃執行。', 18, False, DARK)]],
             algn='just')
    # 版面同完整簡報:頁碼右上、無頁尾
    num = find(s, 'TextBox 7')
    one_line(num, '12')
    place(num, x=31.6, y=0.3)
    for shp in list(s.shapes):
        if shp.name == 'TextBox 6':
            remove(shp)
    s.notes_slide.notes_text_frame.text = (
        '(約20秒)\n陳玨霖老師是本公司的品牌文化顧問，宜蘭縣書法學會創會理事長，在宜蘭教書法將近五十年，'
        '2013年在北京國際文化藝術交流大會書法現場表演得到特別金獎。「舞荳」兩個字就是老師親筆題的。'
        '老師是公司的顧問，本案的文化內容由協力人員規劃、我們執行。')
    prs.save(dst)


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
