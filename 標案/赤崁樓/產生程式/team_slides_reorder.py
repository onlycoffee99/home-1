# -*- coding: utf-8 -*-
"""赤崁樓評選簡報「貳、團隊組織」重排(115/10/12)
用法:python3 -I team_slides_reorder.py 輸入.pptx 輸出.pptx
順序改為:鑫和洋行 → 實績 → 計畫主持人 → 組織架構 → 協力人員(三人合一頁);
原三位協力人員頁移到最後當備用頁(答詢用)。備忘稿一併更新。
"""
import copy
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from team_slides_fill import (BROWN, DARK, RED, T, clone, find, place, remove,  # noqa: E402
                              set_text)
from pptx import Presentation  # noqa: E402

R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'

PEOPLE = [
    ('蔡錦佳', '手作課程顧問',
     ['東華印刷局 計畫主持人', '雲林科大視覺傳達設計碩士', '主持寺廟碑碣、百年老校等文物普查'],
     '規劃手作體驗課程，提供教材設計建議'),
    ('黃荻昌', '文化顧問(鐵道文化)',
     ['長榮大學副教授', '專長鐵道工程與文化資產', '「灣顏號鐵枝路」影音頻道主持人'],
     '規劃「府城交通與城市記憶」系列文化沙龍內容'),
    ('徐鈺清', '臺灣原生物產顧問',
     ['台灣瑰寶牛樟芝生技 負責人', '熟悉原生物產原料來源與品質管理'],
     '協助伴手禮與茶飲原料開發，訂定泡茶用水標準'),
]

NOTES = {
    'pi': '(約20秒)\n我是鄭玉屏，鑫和洋行營運長，也是咖啡烘焙師。剛才那些實績都是我親手做的：門市SOP、人員訓練、菜單和品項開發。'
          '在公部門的場域經營，我很清楚要做到三件事：照規定、留紀錄、先報備。'
          '承攬本案的原因很單純，「把餐飲和文化變成生活產品」是我們一直在做的事，赤崁樓是最適合的地方。',
    'org': '(約15秒)\n本案由我統籌，下面分活動推廣、行銷、餐飲、行政事務四組。另外有三位顧問以個人名義協助，下一頁介紹。',
    'team': '(約20秒)\n三位顧問各補一塊專業：蔡錦佳老師是東華印刷局計畫主持人，規劃手作體驗課程；'
            '黃荻昌老師是長榮大學副教授，規劃「府城交通與城市記憶」文化沙龍；'
            '徐鈺清是台灣瑰寶牛樟芝生技負責人，協助伴手禮、茶飲原料和泡茶用水標準。'
            '顧問負責規劃指導，現場一律由我們的人執行。',
}


def duplicate(prs, src):
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):
        remove(shp)
    for shp in src.shapes:
        el = copy.deepcopy(shp._element)
        for b in el.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}blip'):
            rid = b.get('{%s}embed' % R_NS)
            part = src.part.related_part(rid)
            b.set('{%s}embed' % R_NS, new.part.relate_to(part, R_NS + '/image'))
        new.shapes._spTree.append(el)
    return new


def build_team(prs, base, pi):
    s = duplicate(prs, base)
    keep = {'Rectangle 2', 'TextBox 4', 'Connector 5', 'TextBox 7', 'TextBox 6', '文字方塊 1', '直線接點 12', '文字方塊 17'}
    for shp in list(s.shapes):
        if shp.name not in keep:
            remove(shp)
    name = find(s, '文字方塊 1')
    place(name, w=5.8)
    name.text_frame.paragraphs[0].runs[0].text = '協力人員'
    for r in name.text_frame.paragraphs[0].runs[1:]:
        r.text = ''
    place(find(s, '直線接點 12'), x=7.4)
    role = find(s, '文字方塊 17')
    place(role, x=7.6, w=25)
    role.text_frame.paragraphs[0].runs[0].text = '以個人名義參與｜顧問規劃指導，現場由本團隊執行'
    for r in role.text_frame.paragraphs[0].runs[1:]:
        r.text = ''

    boxbg = find(pi, 'Rounded Rectangle 9', 3)
    bar = find(pi, 'Rectangle 10', 2)
    txt = find(pi, 'TextBox 11')
    for i, (nm, title, cv, duty) in enumerate(PEOPLE):
        x = 1.1 + i * 10.85
        clone(boxbg, s, x=x, y=5.0, w=10.3, h=12.3)
        clone(bar, s, x=x, y=5.0, w=10.3, h=0.25)
        head = clone(txt, s, x=x + 0.4, y=5.5, w=9.5, h=2.4)
        set_text(head, [[T(nm, 28, True, DARK)], [T(title, 18, True, RED)]])
        body = clone(txt, s, x=x + 0.4, y=8.3, w=9.5, h=4.6)
        set_text(body, [[T('■ ', 10, False, RED), T(c, 15, False, BROWN)] for c in cv], line=120)
        line = clone(bar, s, x=x + 0.4, y=13.2, w=9.5, h=0.05)
        duty_box = clone(txt, s, x=x + 0.4, y=13.5, w=9.5, h=3.6)
        set_text(duty_box, [[T('本案負責', 15, True, RED)], [T(duty, 16, False, DARK)]], line=115)
    return s


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def set_pageno(slide, label):
    runs = find(slide, 'TextBox 7').text_frame.paragraphs[0].runs
    runs[0].text = label
    for r in runs[1:]:
        r.text = ''


def main(src, dst):
    prs = Presentation(src)
    S = list(prs.slides)
    company, record, org, pi, cai, huang, xu = S
    team = build_team(prs, cai, pi)
    order = [company, record, pi, org, team, cai, huang, xu]
    labels = ['04', '05', '06', '07', '08', '備1', '備2', '備3']
    sld = prs.slides._sldIdLst
    ids = {s.slide_id: e for s in prs.slides for e in sld if int(e.get('id')) == s.slide_id}
    for s in order:
        el = ids[s.slide_id]
        sld.remove(el)
        sld.append(el)
    for s, lb in zip(order, labels):
        set_pageno(s, lb)
    set_notes(pi, NOTES['pi'])
    set_notes(org, NOTES['org'])
    set_notes(team, NOTES['team'])
    for s in (cai, huang, xu):
        old = s.notes_slide.notes_text_frame.text
        if not old.startswith('【備用頁】'):
            set_notes(s, '【備用頁】委員問到顧問資歷時才翻出。\n' + old)
    prs.save(dst)


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
