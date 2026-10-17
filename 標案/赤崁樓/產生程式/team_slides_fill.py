# -*- coding: utf-8 -*-
"""赤崁樓評選簡報「貳、團隊組織」7 頁補完(115/10/11)
用法:python3 -I team_slides_fill.py 輸入.pptx 輸出.pptx
- 以老闆手改版 pptx 為底,保留原版型、字型與圖示,只補內容
- 協力人員資料來源:資歷庫/協力夥伴.md、人員_徐鈺清.md、content.py(送件版服務建議書)
"""
import copy
import sys

from lxml import etree
from pptx import Presentation
from pptx.util import Cm, Emu

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
FONT = '微軟正黑體'
BROWN = '5E4E40'
DARK = '3B2A20'
RED = 'B94A37'
FOOTER = '鑫和洋行｜赤崁樓遊客中心2樓部分空間委外經營管理案'


def q(tag):
    return '{%s}%s' % (A, tag)


def find(slide, name, nth=0):
    hits = [s for s in slide.shapes if s.name == name]
    return hits[nth]


def next_id(slide):
    return max(int(e.get('id')) for e in slide._element.iter('{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr')) + 1


def clone(src, slide, x=None, y=None, w=None, h=None):
    el = copy.deepcopy(src._element)
    for c in el.iter('{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr'):
        c.set('id', str(next_id(slide)))
        for ext in c.findall(q('extLst')):
            c.remove(ext)
    # 圖片關聯:來源與目標不同頁時要搬 image part
    blips = el.findall('.//' + q('blip'))
    for b in blips:
        rid = b.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
        src_part = src.part.related_part(rid)
        if src.part is not slide.part:
            new_rid = slide.part.relate_to(src_part, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/image')
            b.set('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed', new_rid)
    slide.shapes._spTree.append(el)
    shp = slide.shapes[-1]
    place(shp, x, y, w, h)
    return shp


def place(shp, x=None, y=None, w=None, h=None):
    if x is not None:
        shp.left = Cm(x)
    if y is not None:
        shp.top = Cm(y)
    if w is not None:
        shp.width = Cm(w)
    if h is not None:
        shp.height = Cm(h)


def rpr(sz, bold=False, color=BROWN):
    r = etree.SubElement(etree.Element(q('r')), q('rPr'))
    r.set('lang', 'zh-TW')
    r.set('altLang', 'en-US')
    r.set('sz', str(int(sz * 100)))
    r.set('b', '1' if bold else '0')
    r.set('dirty', '0')
    fill = etree.SubElement(r, q('solidFill'))
    etree.SubElement(fill, q('srgbClr')).set('val', color)
    for t in ('latin', 'ea', 'cs'):
        etree.SubElement(r, q(t)).set('typeface', FONT)
    return r


def set_text(shp, paras, wrap=True, algn='l', spc=None, line=None):
    """paras: [[(text, sz, bold, color), ...], ...]"""
    tf = shp.text_frame
    body = tf._txBody
    bodyPr = body.find(q('bodyPr'))
    bodyPr.set('wrap', 'square' if wrap else 'none')
    for p in body.findall(q('p')):
        body.remove(p)
    for runs in paras:
        p = etree.SubElement(body, q('p'))
        pPr = etree.SubElement(p, q('pPr'))
        pPr.set('algn', algn)
        if line:
            ls = etree.SubElement(pPr, q('lnSpc'))
            etree.SubElement(ls, q('spcPct')).set('val', str(int(line * 1000)))
        if spc:
            sb = etree.SubElement(pPr, q('spcBef'))
            etree.SubElement(sb, q('spcPts')).set('val', str(int(spc * 100)))
        etree.SubElement(pPr, q('buNone'))
        for text, sz, bold, color in runs:
            r = etree.SubElement(p, q('r'))
            r.append(rpr(sz, bold, color))
            etree.SubElement(r, q('t')).text = text


def T(text, sz=16, bold=False, color=BROWN):
    return (text, sz, bold, color)


def bullets(items, sz=16):
    out = []
    for it in items:
        if isinstance(it, tuple):
            head, rest = it
            out.append([T('■ ', sz - 6, False, RED), T(head, sz, True, BROWN), T(rest, sz, False, BROWN)])
        else:
            out.append([T('■ ', sz - 6, False, RED), T(it, sz, False, BROWN)])
    return out


def bar_join(items, sz=22):
    runs = []
    for i, it in enumerate(items):
        if i:
            runs.append(T('∣', sz + 4))
        runs.append(T(it, sz))
    return [runs]


def remove(shp):
    shp._element.getparent().remove(shp._element)


def ensure_footer(slide, tpl, pageno):
    names = [s.name for s in slide.shapes]
    if 'TextBox 6' not in names:
        clone(tpl, slide)
    num = find(slide, 'TextBox 7')
    r = num.text_frame.paragraphs[0].runs
    r[0].text = '%02d' % pageno
    for extra in r[1:]:
        extra.text = ''


def section(slide, bar_tpl, lbl_tpl, y, label):
    clone(bar_tpl, slide, y=y + 0.1)
    lb = clone(lbl_tpl, slide, y=y)
    lb.text_frame.paragraphs[0].runs[0].text = label
    for extra in lb.text_frame.paragraphs[0].runs[1:]:
        extra.text = ''
    return lb


def main(src, dst):
    prs = Presentation(src)
    S = prs.slides
    s4 = S[3]
    tpl = {
        'footer': find(s4, 'TextBox 6'),
        'subtitle': find(s4, '文字方塊 21'),
        'bar': find(s4, 'Rectangle 10', 0),
        'label': find(s4, '文字方塊 26'),
        'big': find(s4, '文字方塊 79'),
        'boxbg': find(s4, 'Rounded Rectangle 9', 3),   # 承攬緣由底框
        'boxbar': find(s4, 'Rectangle 10', 2),
        'boxtxt': find(s4, 'TextBox 11'),
    }
    # 第 2 頁圖示(定位、門市、標籤、人群、城樓)
    s2 = S[1]
    icon2 = {'pin': find(s2, '圖片 31'), 'store': find(s2, '圖片 39'), 'tag': find(s2, '圖片 45'),
             'people': find(s2, '圖片 37'), 'temple': find(s2, '圖片 43')}

    # ---------- 第 1 頁:鑫和洋行 投標主體 ----------
    s1 = S[0]
    sub = clone(tpl['subtitle'], s1, x=1.0, y=4.4, w=31)
    set_text(sub, [[T('民國107年成立 •負責人 張哲源 •主理人 鄭玉屏(本案計畫主持人)', 20)]])
    # 實體店面:兩館改上下排
    y = 5.5
    lbl = find(s1, '文字方塊 26')
    place(lbl, y=y)
    place(find(s1, 'Rectangle 10'), y=y + 0.1)
    st1 = find(s1, '文字方塊 39')
    place(st1, x=1.1, y=y + 1.3, w=21.5, h=1.0)
    set_text(st1, [[T('本館', 18, True, RED), T('｜宜蘭礁溪溫泉公園「時間到咖啡館－公園裡的咖啡館」', 18)],
                   [T('二館', 18, True, RED), T('｜臺中市第二市場停車場附屬空間(台中館，115/9/19開幕)', 18)]], line=125)
    remove(find(s1, '文字方塊 42'))
    section(s1, tpl['bar'], tpl['label'], 9.2, '主要業務')
    b1 = clone(tpl['big'], s1, x=1.0, y=10.4)
    set_text(b1, bar_join(['咖啡豆選品與烘焙', '製茶與飲品服務', '農產品銷售', '國際貿易'], 20), wrap=False)
    section(s1, tpl['bar'], tpl['label'], 11.9, '自有品牌')
    b2 = clone(tpl['big'], s1, x=1.0, y=13.1)
    set_text(b2, [[T('「時間到」', 20, True), T('咖啡館／茶館', 20), T('∣', 24),
                   T('「舞荳」', 20, True), T('商標註冊第02507640號', 20)]], wrap=False)
    bg = clone(tpl['boxbg'], s1, x=1.1, y=14.6, w=21.8, h=3.0)
    clone(tpl['boxbar'], s1, x=1.1, y=14.6, w=0.2, h=3.0)
    bt = clone(tpl['boxtxt'], s1, x=1.4, y=14.7, w=21.4, h=2.8)
    set_text(bt, [[T('公司簡介：', 18, True, RED),
                   T('以咖啡豆選品、烘焙及咖啡館經營為基礎，創立「時間到」品牌，自宜蘭礁溪發跡，'
                     '透過產學合作及在地特產開發特色商品，現已拓展至臺中市中區，累積公共場域進駐與跨區門市經營經驗。', 18, False, DARK)]],
             algn='just')
    # 右側:數字卡(礁溪照片第 2 頁已有,不重複放)
    for pic_name in ('Picture 11', 'TextBox 12'):  # 第 2 頁版面外殘留物件
        remove(find(s2, pic_name))
    stats = [('107年', '成立至今'), ('2館', '跨區門市(宜蘭／臺中)'), ('02507640', '「舞荳」註冊商標')]
    for i, (big, small) in enumerate(stats):
        y0 = 5.6 + i * 3.0
        clone(tpl['boxbg'], s1, x=23.6, y=y0, w=9.6, h=2.6)
        tb = clone(tpl['boxtxt'], s1, x=23.8, y=y0 + 0.2, w=9.2, h=2.3)
        set_text(tb, [[T(big, 26, True, RED)], [T(small, 14, False, DARK)]], algn='ctr')
    tagbox = clone(tpl['boxbg'], s1, x=23.6, y=14.6, w=9.6, h=3.0)
    tb = clone(tpl['boxtxt'], s1, x=23.8, y=14.75, w=9.2, h=2.8)
    set_text(tb, [[T('本案由鑫和洋行單獨投標', 18, True, RED)],
                  [T('負全部履約責任；協力人員以個人名義提供專業協助', 14, False, DARK)]], algn='ctr')

    # ---------- 第 2 頁:圖示說明 ----------
    labels = [('文字方塊 41', 0.9, '公共場域|進駐經營'), ('文字方塊 47', 5.9, '跨區門市|經營'),
              ('文字方塊 49', 11.1, '自有品牌|與商品')]
    lbl_src = find(s2, '文字方塊 41')
    extra = [(16.7, '產學合作|與推廣教學'), (21.3, '文化內容|轉譯')]
    for x, txt in extra:
        n = clone(lbl_src, s2, x=x, y=15.7)
        labels.append((None, x, txt, n))
    for item in labels:
        shp = item[3] if len(item) == 4 else find(s2, item[0])
        place(shp, x=item[1], y=15.75, w=4.4)
        a, b = item[2].split('|')
        set_text(shp, [[T(a, 14)], [T(b, 14)]], algn='ctr')

    # ---------- 第 5~7 頁:協力人員 ----------
    people = [
        dict(sub='東華印刷局資訊科技企業社 計畫主持人 •視覺傳達設計 •文物普查',
             cards=[('people', '課程規劃', '印刷手作、教材設計'),
                    ('doc', '視覺傳達', '展示設計、出版設計'),
                    ('temple', '文物普查', '田野調查、口述歷史')],
             cv=[('現職　', '東華印刷局資訊科技企業社計畫主持人(東華印刷局1977年開業，深耕屏東近半世紀)'),
                 ('學歷　', '國立雲林科技大學視覺傳達設計系碩士'),
                 ('經歷　', '國際日本文化研究中心共同研究員(2017～2022年)'),
                 ('實績　', '主持屏東縣寺廟與碑碣、百年老校、泰武鄉原住民族文物普查及台電文化資產清查等計畫')],
             role='手作課程顧問', duty='規劃印刷手作體驗課程，提供教材設計建議；現場課程由本團隊人員帶領。'),
        dict(sub='長榮大學副教授 •鐵道工程與文化資產 •「灣顏號鐵枝路」頻道主持人',
             cards=[('temple', '鐵道文資', '交通建設與城市發展'),
                    ('people', '影音推廣', '灣顏號鐵枝路頻道'),
                    ('cup', '文化沙龍', '府城交通與城市記憶')],
             cv=[('現職　', '長榮大學副教授'),
                 ('專長　', '鐵道工程與文化資產、TOD大眾運輸導向發展規劃、土木工程、永續規劃與設計'),
                 ('經歷　', '2026年升任土木學會歷史文化相關職務；長期推動「鐵道文化經典之旅」'),
                 ('推廣　', '主持「灣顏號鐵枝路-黃教授」影音頻道，製作《牡丹社事件後台灣鐵道發展經緯》等影片')],
             role='文化顧問', duty='規劃「府城交通與城市記憶」系列文化沙龍內容；活動由本團隊人員執行。'),
        dict(sub='台灣瑰寶牛樟芝生物科技有限公司 負責人 •臺灣原生物產',
             cards=[('cup', '茶飲原料', '原生物產、萃茶用水'),
                    ('doc', '品質管理', '原料來源、開發流程'),
                    ('tag', '商品開發', '文創伴手禮')],
             cv=[('現職　', '台灣瑰寶牛樟芝生物科技有限公司負責人'),
                 ('專長　', '臺灣原生牛樟芝相關產品之研發與銷售'),
                 ('能力　', '熟悉臺灣原生物種之原料來源、品質管理及產品開發流程')],
             role='臺灣原生物產顧問',
             duty='協助文創伴手禮與茶飲原料開發；研討泡茶用水(水質、過濾、煮水溫度)，訂定標準用水及萃茶參數。商品標示依食品安全衛生管理法規辦理。'),
    ]
    for idx, d in enumerate(people):
        s = S[4 + idx]
        sub = clone(tpl['subtitle'], s, x=1.0, y=4.4, w=31)
        set_text(sub, [[T(d['sub'], 20)]])
        section(s, tpl['bar'], tpl['label'], 5.5, '專長領域')
        # 版面外的三張能力卡移進版面
        cards_x = [1.4, 9.6, 17.7]
        dy = Cm(7.0) - Cm(19.2)
        icons = {'people': find(s, '圖片 36'), 'cup': find(s, '圖片 38'), 'doc': find(s, '圖片 67')}
        groups = [('Rounded Rectangle 9', 1), ('Rounded Rectangle 9', 2), ('Rounded Rectangle 9', 0)]
        texts = [('文字方塊 40', '文字方塊 45', '直線接點 41'), ('文字方塊 70', '文字方塊 72', '直線接點 71'),
                 ('文字方塊 75', '文字方塊 77', '直線接點 76')]
        rects = sorted([sh for sh in s.shapes if sh.name == 'Rounded Rectangle 9'], key=lambda z: z.left)
        icon_slots = sorted(icons.values(), key=lambda z: z.left)
        slot_pos = [(z.left, z.top, z.width, z.height) for z in icon_slots]
        for sh in list(rects) + [find(s, n) for trio in texts for n in trio]:
            sh.top = sh.top + dy
        used = set()
        for i, (ic, head, desc) in enumerate(d['cards']):
            l, t, w, h = slot_pos[i]
            if ic in icons and ic not in used:
                shp = icons[ic]
                used.add(ic)
                shp.left, shp.top = l, t + dy
                s.shapes._spTree.append(shp._element)
            else:
                shp = clone(icon2[ic], s)
                shp.left, shp.top, shp.width, shp.height = l, t + dy, w, h
            s.shapes._spTree.append(shp._element)  # 圖示置頂,免被卡片底框蓋住
            th, td, _ = texts[i]
            find(s, th).text_frame.paragraphs[0].runs[0].text = head
            set_text(find(s, td), [[T(desc, 16)]], wrap=False, algn='just')
        for k, shp in icons.items():
            if k not in used:
                remove(shp)
        section(s, tpl['bar'], tpl['label'], 9.6, '學經歷')
        cv = clone(tpl['big'], s, x=1.0, y=10.8, w=31.5, h=3.6)
        set_text(cv, bullets(d['cv'], 17), line=110)
        place(cv, w=31.5)
        clone(tpl['boxbg'], s, x=1.1, y=14.75, w=31.6, h=2.8)
        clone(tpl['boxbar'], s, x=1.1, y=14.75, w=0.2, h=2.8)
        bt = clone(tpl['boxtxt'], s, x=1.4, y=14.85, w=31.1, h=2.6)
        set_text(bt, [[T('本案分工｜' + d['role'] + '：', 20, True, RED), T(d['duty'], 20, False, DARK)]], algn='just')

    # 頁尾與頁碼(04 起連號)
    for i, s in enumerate(S):
        ensure_footer(s, tpl['footer'], 4 + i)
    prs.save(dst)


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
