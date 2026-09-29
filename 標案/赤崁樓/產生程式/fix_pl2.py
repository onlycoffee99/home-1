import sys, copy, docx
from docx.oxml.ns import qn
SRC, OUT = sys.argv[1:3]
d = docx.Document(SRC)
FB,RET,EVT=9328800,720000,360000; RAMP=(0.8,0.92,0.97); PK=0.05; FIX,TH=550000,5400000
n=lambda x: f'{round(x):,}'
Y=[]
for i,r in enumerate(RAMP):
    fb,ret,evt=FB*r,RET*r,EVT*r; rev=fb+ret+evt
    c={'餐飲原物料成本(餐飲營收32%)':fb*0.32,'包材耗材(餐飲營收5%)':fb*PK,'商品進貨成本(商品營收55%)':ret*0.55,
       '人事費用(含勞健保，年調3%)':2880000*1.03**i,'固定權利金':FIX,'變動權利金(超過540萬部分2%)':max(0,rev-TH)*0.02,
       '水電瓦斯通訊(年增3%)':420000*1.03**i,'行銷推廣費(營業額3%)':rev*0.03,'支付手續費(營業額1.5%)':rev*0.015,
       '保險、清潔及設備維護':180000,'耗材及雜支(營業額2%)':rev*0.02,'開辦裝修及設備攤提(3年)':566667}
    Y.append(dict(fb=fb,ret=ret,evt=evt,rev=rev,c=c,tot=sum(round(v) for v in c.values())))
def set_text(el, text):
    ts=list(el.iter(qn('w:t'))); ts[0].text=text
    ts[0].set('{http://www.w3.org/XML/1998/namespace}space','preserve')
    for t in ts[1:]: t.text=''
def ucells(row):
    out=[]
    for c in row.cells:
        if c._tc not in [x._tc for x in out]: out.append(c)
    return out
def fill(row, label, vals):
    cs=ucells(row)
    if label is not None: set_text(cs[0]._tc, label)
    if len(cs)==2:  # 三格合併的列,先拆開不可行 → 直接改值不同時換成獨立列
        raise SystemExit('merged row: '+cs[0].text)
    for c,v in zip(cs[1:],vals): set_text(c._tc, v)
# ---- 表15 營收預估 ----
t15=next(t for t in d.tables if t.rows[0].cells[0].text.strip()=='項目' and len(t.rows)==5)
fill(t15.rows[1],None,[n(y['fb']) for y in Y]); fill(t15.rows[2],None,[n(y['ret']) for y in Y])
fill(t15.rows[3],None,[n(y['evt']) for y in Y]); fill(t15.rows[4],None,[n(y['rev']) for y in Y])
# ---- 表16 損益 ----
t16=next(t for t in d.tables if any('固定權利金' in c.text for r in t.rows for c in r.cells))
tbl=t16._tbl
def row_of(prefix):
    return next(r for r in t16.rows if r.cells[0].text.strip().startswith(prefix))
# 固定權利金列若合併,換成可填三格的列(複製原料成本列格式)
template=row_of('餐飲原物料')
fr=row_of('固定權利金')
if len(ucells(fr))==2:
    nr=copy.deepcopy(template._tr); fr._tr.addprevious(nr); tbl.remove(fr._tr)
wr=row_of('水電'); 
if len(ucells(wr))==2:
    nr=copy.deepcopy(template._tr); wr._tr.addprevious(nr); tbl.remove(wr._tr); set_text(nr, '　水電瓦斯及通訊費')
for pre in ('保險','開辦'):
    rr=row_of(pre)
    if len(ucells(rr))==2:
        nr=copy.deepcopy(template._tr); lab=rr.cells[0].text; rr._tr.addprevious(nr); tbl.remove(rr._tr)
# 重新依序寫入成本列
order=list(Y[0]['c'].keys())
start=row_of('營業成本及費用'); end=row_of('成本及費用合計')
# 移除舊成本列
r=start._tr.getnext()
while r is not end._tr:
    nx=r.getnext(); tbl.remove(r); r=nx
prev=start._tr
for k in order:
    nr=copy.deepcopy(template._tr); prev.addnext(nr); prev=nr
from docx.table import _Row
rows=list(t16.rows)
si=[i for i,r in enumerate(rows) if r._tr is start._tr][0]
for j,k in enumerate(order):
    rw=rows[si+1+j]; cs=ucells(rw)
    set_text(cs[0]._tc,'　'+k)
    for c,y in zip(cs[1:],Y): set_text(c._tc, n(y['c'][k]))
def setrow(prefix, vals):
    rw=row_of(prefix); cs=ucells(rw)
    for c,v in zip(cs[1:],vals): set_text(c._tc,v)
for r in t16.rows:
    lab=r.cells[0].text.strip(); cs=ucells(r)
    if len(cs)!=4: continue
    key={'餐飲收入':'fb','文創商品及伴手禮收入':'ret','文化沙龍、包場及活動收入':'evt','營業收入合計':'rev'}
    for kk,vv in key.items():
        if lab.startswith(kk):
            for c,y in zip(cs[1:],Y): set_text(c._tc,n(y[vv]))
setrow('成本及費用合計',[n(y['tot']) for y in Y])
pro=[y['rev']-y['tot'] for y in Y]
setrow('稅前損益',[n(p) for p in pro]); setrow('稅前淨利率',[f'{p/y["rev"]*100:.1f}%' for p,y in zip(pro,Y)])
# ---- 文字 ----
y2=Y[1]
F=2880000*1.03+FIX-TH*0.02+420000*1.03+180000+566667
vr=(y2['fb']*(0.32+PK)+y2['ret']*0.55)/y2['rev']+0.03+0.015+0.02+0.02
BEP=F/(1-vr); BEPM=BEP/12; cust=(BEPM-(RET+EVT)*0.92/12)/26/115
for p in d.paragraphs:
    if p.text.startswith('營收預估以每月營業26日'):
        set_text(p._p,'營收預估以每月營業26日、平均每日來客數260人次(平日約200人次、假日約380人次)、平均客單價115元為穩定營運水準，另計文創商品及活動收入；考量古蹟場域客流受季節與天候影響，第一年市場導入期以穩定水準之80%估算，第二年92%，第三年97%，保留來客波動之緩衝。本預估以來客數達到上述基本面為前提，實際營收將隨來客數增減而變動。')
    if p.text.startswith('經試算，本案年固定成本約'):
        set_text(p._p,f'以第二年為基準，年固定成本約{n(F)}元，變動成本率約{vr*100:.1f}%，損益兩平年營業額約{n(BEP)}元，約每日{cust:.0f}人次。換言之，每日平均來客須達約{cust:.0f}人次方能損益兩平，達約{round(260*RAMP[1])}人次方能實現第二年之預估獲利；第一年導入期約損益兩平，第二、三年稅前淨利率分別約{pro[1]/Y[1]["rev"]*100:.1f}%及{pro[2]/Y[2]["rev"]*100:.1f}%。另每年自盈餘提撥營業額1%作為設備汰換準備金。')
d.save(OUT)
print([round(y['rev']) for y in Y],[y['tot'] for y in Y],[round(p) for p in pro],round(BEP),round(cust))
