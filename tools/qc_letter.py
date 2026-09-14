#!/usr/bin/env python3
# 輕車悠遊公司函產生器:依 縣府案件/範本/輕車函文範本.pdf 版面輸出 Word。
# 用法: python3 tools/qc_letter.py <spec.json> <輸出.docx>
# spec 欄位: date_roc(省略則自動帶今日)、seq(當日序號,預設1)、subject、
#   sections([{title,paras[]}] 或字串)、attachments_line、copies_main、copies_cc、sign(預設true)
import json, sys, datetime
from docx import Document
from docx.shared import Pt, Cm, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

KAI = '標楷體'
TPL_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '縣府案件', '範本')
SIG_IMG = os.path.join(TPL_DIR, '董事長簽名.jpg')


def set_font(run, size, bold=False):
    run.font.name = KAI
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), KAI)


def para(doc, text, size=12, bold=False, align=None, left=None, hang=None, before=0, after=0, line=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = Pt(line)
    if align is not None:
        pf.alignment = align
    if left is not None:
        pf.left_indent = Cm(left)
    if hang is not None:
        pf.first_line_indent = Cm(-hang)
    r = p.add_run(text)
    set_font(r, size, bold)
    return p


def roc_today():
    t = datetime.date.today()
    return t.year - 1911, t.month, t.day


def add_footer(doc):
    sec = doc.sections[0]
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    def fld(code):
        r = p.add_run()
        for el, attr in (('w:fldChar', {'w:fldCharType': 'begin'}), ('w:instrText', None), ('w:fldChar', {'w:fldCharType': 'end'})):
            e = OxmlElement(el)
            if attr:
                for k, v in attr.items():
                    e.set(qn(k), v)
            else:
                e.set(qn('xml:space'), 'preserve')
                e.text = code
            r._element.append(e)
        set_font(r, 10)
    r = p.add_run('第 '); set_font(r, 10)
    fld('PAGE')
    r = p.add_run(' 頁 共 '); set_font(r, 10)
    fld('NUMPAGES')
    r = p.add_run(' 頁'); set_font(r, 10)


def build(spec, out):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Mm(297), Mm(210)
    sec.top_margin, sec.bottom_margin = Cm(1.5), Cm(2.0)
    sec.left_margin, sec.right_margin = Cm(2.5), Cm(2.5)

    # 正本框(左上,1x1 表格)
    tb = doc.add_table(rows=1, cols=1)
    tb.alignment = WD_TABLE_ALIGNMENT.LEFT
    tb.autofit = False
    tblPr = tb._tbl.tblPr
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), str(int(2.6 * 567))); tblW.set(qn('w:type'), 'dxa')
    tblPr.append(tblW)
    layout = OxmlElement('w:tblLayout'); layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)
    tb.columns[0].width = Cm(2.6)
    cell = tb.cell(0, 0)
    cell.width = Cm(2.6)
    cp = cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cp.add_run('正　本')
    set_font(cr, 12)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{side}')
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '8'); e.set(qn('w:color'), '000000')
        borders.append(e)
    tcPr.append(borders)

    # 標題
    para(doc, '輕車悠遊股份有限公司　函', 20, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=6)

    # 右側基本資料區塊
    for t in ('地　　址:262宜蘭縣礁溪鄉公園路16號(營運地點)',
              '承 辦 人:鄭玉屏(手機:0975-328-779)',
              '電　　話:03 987 6416',
              '電子信箱:onlycoffee99@gmail.com'):
        para(doc, t, 11, left=7.8, line=14)

    para(doc, '', 8)
    para(doc, '260', 11)
    para(doc, '宜蘭市凱旋里三鄰縣政北路一號', 12)
    para(doc, '', 6)
    para(doc, '受文者:宜蘭縣政府', 15, bold=True, after=6)

    y, m, d = roc_today()
    if spec.get('date_roc'):
        y, m, d = spec['date_roc']
    seq = spec.get('seq', 1)
    para(doc, f'發文日期:中華民國{y}年{m}月{d}日', 12)
    para(doc, f'發文字號:{y}北輕字第{y}{m:02d}{d:02d}-{seq}號', 12)
    para(doc, '速　　別:普通件', 12)
    para(doc, '密等及解密條件或保密期限:普通', 12)
    para(doc, f"附　　件:{spec.get('attachments_line', '略。')}", 12, after=8)

    # 主旨
    para(doc, f"主旨:{spec['subject']}", 14, left=1.7, hang=1.7, after=8, line=22)

    # 說明
    para(doc, '說明:', 14)
    for i, s in enumerate(spec.get('sections', [])):
        num = '一二三四五六七八九十'[i]
        if isinstance(s, str):
            para(doc, f'{num}、{s}', 14, left=1.2, hang=0.65, after=4, line=22)
        else:
            para(doc, f"{num}、{s['title']}" if s.get('title') else f'{num}、', 14, left=1.2, hang=0.65, after=2, line=22)
            for sub in s.get('paras', []):
                para(doc, sub, 14, left=2.4, hang=1.1, after=4, line=22)

    para(doc, '', 6)
    para(doc, f"正本:{spec.get('copies_main', '宜蘭縣政府')}", 12)
    para(doc, f"副本:{spec.get('copies_cc', '輕車悠遊股份有限公司')}", 12, after=12)

    if spec.get('sign', True) and os.path.exists(SIG_IMG):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.add_run().add_picture(SIG_IMG, width=Cm(10))

    add_footer(doc)
    doc.save(out)
    print('WROTE', out)


if __name__ == '__main__':
    with open(sys.argv[1], encoding='utf-8') as f:
        spec = json.load(f)
    build(spec, sys.argv[2])
